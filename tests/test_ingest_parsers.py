"""Tests for the format parsers.

They assert on ``ParsedUnit`` — the parser's own output — never on
``Document``. That is the payoff of keeping the two apart: no parser test
has to invent catalog metadata to satisfy a schema.
"""

from __future__ import annotations

import io
import json

import pymupdf
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.ingest.parsers import DEFAULT_PARSERS


def _pdf_bytes(pages: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(56, 56, 539, 700), text, fontsize=10)
    raw = doc.tobytes()
    doc.close()
    return raw


# ----------------------------------------------------------------------
# JSON budgets
# ----------------------------------------------------------------------
def test_budget_json_renders_structured_markdown_not_a_json_dump() -> None:
    """A json.dumps blob mixes technical keys with semantic values and
    makes every record's vector look alike."""
    raw = json.dumps({
        "budget_id": "BUDGET-2024-0315", "client_name": "Acme Corp",
        "project_type": "CRM", "total_amount": 80000, "currency": "EUR",
        "phases": [{"name": "Diseño", "hours": 40, "cost_eur": 2500}],
    }).encode()

    unit = DEFAULT_PARSERS["json"].parse(raw, "b.json")[0]

    assert unit.content.startswith("# Presupuesto BUDGET-2024-0315")
    assert "Cliente: Acme Corp." in unit.content
    assert "| Diseño | 40 | 2500 |" in unit.content
    assert "budget_id" not in unit.content  # the key itself never ships
    assert unit.document_id == "BUDGET-2024-0315"


def test_budget_parser_does_not_clean_its_own_records() -> None:
    """A parser that quietly fixed a date format would hide the very
    signal the audit exists to surface."""
    raw = json.dumps({"budget_id": "B1", "signed_at": "15/03/2024", "currency": "€"}).encode()

    records = DEFAULT_PARSERS["json"].to_records(raw, "b.json")

    assert records[0]["signed_at"] == "15/03/2024"
    assert records[0]["currency"] == "€"


# ----------------------------------------------------------------------
# TXT transcripts — the two format eras
# ----------------------------------------------------------------------
def test_tagged_transcript_yields_one_unit_per_turn_with_speaker() -> None:
    raw = (
        "[00:01:12] Ana Ruiz: Necesitamos multi-tenant desde el principio.\n"
        "[00:02:03] Luis Soto: El presupuesto es de 30000 euros.\n"
    ).encode()

    units = DEFAULT_PARSERS["txt"].parse(raw, "2024_kickoff.txt")

    assert [u.extra["speaker"] for u in units] == ["Ana Ruiz", "Luis Soto"]
    assert units[0].extra["timestamp"] == "00:01:12"
    assert all(u.extra["format_era"] == "tagged" for u in units)


def test_legacy_transcript_falls_back_to_paragraphs_and_says_so() -> None:
    """Attribution is unavailable, and recording that is itself a quality
    signal — not something to paper over."""
    raw = "Primer bloque de la reunión.\n\nSegundo bloque del acta.".encode()

    units = DEFAULT_PARSERS["txt"].parse(raw, "2022_acta.txt")

    assert len(units) == 2
    assert all(u.extra["format_era"] == "legacy_untagged" for u in units)
    assert all(u.extra["speaker_attribution"] == "unavailable" for u in units)


def test_a_stray_colon_line_does_not_promote_a_legacy_transcript() -> None:
    raw = (
        "Nota: la reunión se celebró por videollamada.\n\n"
        "El cliente explicó su situación actual con detalle.\n\n"
        "Se acordó revisar el alcance la semana siguiente.\n\n"
        "Quedamos en enviar una propuesta el viernes.\n"
    ).encode()

    units = DEFAULT_PARSERS["txt"].parse(raw, "old.txt")

    assert all(u.extra["format_era"] == "legacy_untagged" for u in units)


def test_an_untagged_line_inside_a_tagged_transcript_is_kept() -> None:
    raw = (
        "[00:01:12] Ana Ruiz: Empezamos.\n"
        "(ruido de fondo)\n"
        "[00:02:03] Luis Soto: Continuamos.\n"
    ).encode()

    units = DEFAULT_PARSERS["txt"].parse(raw, "t.txt")

    assert len(units) == 3
    assert units[1].extra["untagged_line"] is True


# ----------------------------------------------------------------------
# DOCX proposals
# ----------------------------------------------------------------------
def test_docx_splits_into_sections_by_heading() -> None:
    document = DocxDocument()
    document.add_heading("Alcance", level=1)
    document.add_paragraph("El proyecto cubre la plataforma completa.")
    document.add_heading("Cronograma", level=1)
    document.add_paragraph("Doce semanas repartidas en tres fases.")
    buffer = io.BytesIO()
    document.save(buffer)

    units = DEFAULT_PARSERS["docx"].parse(buffer.getvalue(), "propuesta.docx")

    assert [u.section_title for u in units] == ["Alcance", "Cronograma"]
    assert "plataforma completa" in units[0].content


# ----------------------------------------------------------------------
# PDF contracts
# ----------------------------------------------------------------------
def test_pdf_yields_one_unit_per_page_with_page_numbers() -> None:
    raw = _pdf_bytes(["CONTRATO. Cláusula primera del acuerdo marco.",
                      "Cláusula segunda. Condiciones económicas acordadas."])

    units = DEFAULT_PARSERS["pdf"].parse(raw, "contrato.pdf")

    assert [u.page_number for u in units] == [1, 2]
    assert "Cláusula segunda" in units[1].content


def test_an_unreadable_pdf_yields_nothing_rather_than_an_empty_unit() -> None:
    """An empty Document would be embedded as a vector meaning nothing,
    which is worse than an absence the orchestrator can count."""
    units = DEFAULT_PARSERS["pdf"].parse(b"%PDF-1.4 not really a pdf", "broken.pdf")

    assert units == []


# ----------------------------------------------------------------------
# XLSX rate card
# ----------------------------------------------------------------------
def test_xlsx_becomes_a_markdown_table() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tarifas"
    for row in [["Rol", "Seniority", "EUR/h"], ["Backend", "Senior", 62.5]]:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)

    units = DEFAULT_PARSERS["xlsx"].parse(buffer.getvalue(), "tarifas.xlsx")

    assert units[0].section_title == "Tarifas"
    assert "| Rol | Seniority | EUR/h |" in units[0].content
    assert "| Backend | Senior | 62.5 |" in units[0].content


def test_hidden_sheets_are_excluded() -> None:
    """A hidden sheet is hidden for a reason the pipeline cannot know."""
    workbook = Workbook()
    visible = workbook.active
    visible.title = "Visible"
    visible.append(["a", "b"])
    visible.append(["1", "2"])
    hidden = workbook.create_sheet("Oculta")
    hidden.append(["x", "y"])
    hidden.append(["3", "4"])
    hidden.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)

    units = DEFAULT_PARSERS["xlsx"].parse(buffer.getvalue(), "libro.xlsx")

    assert [u.section_title for u in units] == ["Visible"]


def test_every_registered_parser_declares_its_formats() -> None:
    for fmt, parser in DEFAULT_PARSERS.items():
        assert fmt in parser.supported_formats
