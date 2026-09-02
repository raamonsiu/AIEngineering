"""Build the synthetic "enterprise" corpus that the ingestion work is tested against.

Run it with::

    uv run python data/build_corpus.py

Why this exists
---------------
The estimator is fed real client material: budgets exported from an ERP,
meeting transcripts, Word proposals, signed PDFs, a rate card. We cannot put
real client data in the repository, and we cannot test a cleaning pipeline
against data that is already clean. So this script fabricates a corpus that
is *deliberately* as messy as the real one, with every defect written down in
``data/README.md`` so a later test can assert on it instead of guessing.

Design decisions
----------------
**One script, five sources, no parsers.** This module only writes files. It
never reads the corpus back to decide what to write, so the corpus cannot
drift into agreeing with a buggy reader.

**The defect table is the source of truth, not a comment.** ``PLANTED_DEFECTS``
is a single tuple that both *applies* the defect (most entries carry the exact
value to splice into the JSON record) and *renders* the README table. A defect
cannot therefore be documented but not planted, or planted but not documented.

**Byte-for-byte deterministic.** Same bytes on every run, so the corpus can be
committed and a re-run shows an empty diff. That needs three things the
libraries do not give for free: a seeded RNG (never module-level ``random``),
fixed timestamps baked into the documents (never ``datetime.now()``), and two
post-processing steps — PyMuPDF stamps a random ``/ID`` into every PDF trailer
and ZIP containers (docx/xlsx) record the wall clock per entry.

**The rate card is deliberately stale.** It is the one source that must be
*excluded* downstream, and the exclusion rule is "last modified more than N
months ago". That signal lives on the filesystem, not in the bytes, so the
script sets the mtimes explicitly: everything else looks freshly exported,
the rate card looks like January 2024.

**Money is coherent with ``app/context/examples.py``.** Every clean budget
satisfies ``total_amount == hours_estimated * 62.50`` and every phase cost is
its hours at the same rate, because the few-shot examples the model is primed
with use that implicit rate. A corpus on another rate would teach the
retrieval layer a different economics from the prompt layer.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import random
import re
import shutil
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Final

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Global knobs
# --------------------------------------------------------------------------

#: Single seed for every random draw in this module. Changing it reshuffles the
#: filler text of the transcripts but never the planted defects, which are
#: literal values rather than draws.
SEED: Final[int] = 20240115

#: The implicit rate of ESTIMATION_EXAMPLES (2500/40 == 3750/60 == 62.50).
RATE_EUR_PER_HOUR: Final[float] = 62.50

DATA_DIR: Final[Path] = Path(__file__).resolve().parent
CORPUS_DIR: Final[Path] = DATA_DIR / "corpus"

#: Fixed mtimes. "Fresh" is recent enough that a staleness rule lets it through;
#: the rate card sits far enough back that the same rule rejects it.
FRESH_MTIME: Final[datetime] = datetime(2026, 9, 1, 8, 0, 0, tzinfo=timezone.utc)
STALE_MTIME: Final[datetime] = datetime(2024, 1, 18, 9, 30, 0, tzinfo=timezone.utc)

#: Baked into document metadata so containers hash identically across runs.
DOC_TIMESTAMP: Final[datetime] = datetime(2024, 1, 1, 0, 0, 0)

VENDOR_NAME: Final[str] = "Lidrai Soluciones Digitales S.L."
VENDOR_CIF: Final[str] = "B-87654321"
VENDOR_ADDRESS: Final[str] = "Calle Velázquez 45, 3º B, 28001 Madrid"
VENDOR_EMAIL: Final[str] = "proyectos@lidrai.es"
VENDOR_PHONE: Final[str] = "+34 910 22 44 86"

# Defect family labels. Used verbatim in the README table, so a test can group
# rows by this exact string.
F1: Final[str] = "F1 Heterogeneidad de formato"
F2: Final[str] = "F2 Duplicados divergentes"
F3: Final[str] = "F3 Nulos disfrazados"
F4: Final[str] = "F4 Fuera de rango"


# --------------------------------------------------------------------------
# Domain model
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Client:
    """A client of the consultancy, with the PII a transcript would leak."""

    code: str
    name: str  # canonical spelling; variants are planted as defects
    slug: str
    city: str
    contact_name: str
    contact_email: str
    contact_phone: str
    second_attendee: str


CLIENTS: Final[dict[str, Client]] = {
    c.code: c
    for c in (
        Client(
            "CLI-1001",
            "ACME Corp.",
            "acme-corp",
            "Madrid",
            "Lucía Fernández Arriaga",
            "lucia.fernandez@acme-corp.es",
            "+34 915 44 20 18",
            "Tomás Iriarte Blanco",
        ),
        Client(
            "CLI-1002",
            "Textiles Moreno S.L.",
            "textiles-moreno",
            "Valencia",
            "Javier Moreno Ibáñez",
            "javier.moreno@textilesmoreno.es",
            "+34 963 87 11 02",
            "Rocío Benavent Gil",
        ),
        Client(
            "CLI-1003",
            "Grupo Navarro Logística",
            "grupo-navarro",
            "Zaragoza",
            "Sergio Navarro Puig",
            "s.navarro@gruponavarro.com",
            "+34 976 23 55 41",
            "Alicia Used Mainar",
        ),
        Client(
            "CLI-1004",
            "Farmacias Delgado",
            "farmacias-delgado",
            "Sevilla",
            "Marta Delgado Ruiz",
            "marta.delgado@farmaciasdelgado.es",
            "+34 954 02 76 33",
            "Ignacio Pavón Lara",
        ),
        Client(
            "CLI-1005",
            "Bodegas Villanueva",
            "bodegas-villanueva",
            "Logroño",
            "Andrés Villanueva Soto",
            "andres@bodegasvillanueva.es",
            "+34 941 60 18 94",
            "Celia Ochoa Arnedo",
        ),
        Client(
            "CLI-1006",
            "Clínica Sanchís",
            "clinica-sanchis",
            "Valencia",
            "Elena Sanchís Mompó",
            "elena.sanchis@clinicasanchis.es",
            "+34 961 33 48 27",
            "Dr. Vicente Alabau Ferrer",
        ),
        Client(
            "CLI-1007",
            "Inmobiliaria Peñalver",
            "inmobiliaria-penalver",
            "Málaga",
            "Rubén Peñalver Gil",
            "ruben.penalver@penalver-inmo.es",
            "+34 952 71 09 65",
            "Sandra Quirós Maldonado",
        ),
        Client(
            "CLI-1008",
            "Distribuciones Arroyo",
            "distribuciones-arroyo",
            "Valladolid",
            "Noelia Arroyo Cuesta",
            "noelia.arroyo@distribucionesarroyo.com",
            "+34 983 15 62 70",
            "Marcos Zorrilla Peña",
        ),
        Client(
            "CLI-1009",
            "Editorial Marimón",
            "editorial-marimon",
            "Barcelona",
            "Pau Marimón Riera",
            "pau.marimon@editorialmarimon.cat",
            "+34 934 28 90 56",
            "Gemma Estany Vidal",
        ),
        Client(
            "CLI-1010",
            "Talleres Beltrán",
            "talleres-beltran",
            "Burgos",
            "Óscar Beltrán Lázaro",
            "oscar.beltran@talleresbeltran.es",
            "+34 947 44 37 12",
            "Laura Cantero Hoyos",
        ),
    )
}

#: Consultancy-side staff. Fixed list so ``account_manager`` is reproducible.
ACCOUNT_MANAGERS: Final[tuple[str, ...]] = (
    "Patricia Olmedo Rubio",
    "Daniel Quintana Vega",
    "Irene Castaño Abad",
    "Víctor Salazar Prieto",
    "Nuria Bermejo Lastra",
)

TECH_LEADS: Final[tuple[str, ...]] = (
    "Álvaro Cifuentes Nieto",
    "Miriam Escudero Pardo",
    "Jorge Tejedor Arroyo",
)

#: The six project types of ESTIMATION_EXAMPLES, with the Spanish wording the
#: documents use.
PROJECT_LABELS: Final[dict[str, str]] = {
    "inventory_platform": "plataforma de gestión de inventario",
    "mobile_ecommerce": "aplicación móvil de comercio electrónico",
    "analytics_dashboard": "cuadro de mando analítico en tiempo real",
    "crm": "sistema CRM",
    "cms": "gestor de contenidos (CMS)",
    "api_integration": "servicio de integración de APIs",
}

#: Phase breakdowns mirroring the task tables of ESTIMATION_EXAMPLES, expressed
#: as weights so any total of hours can be split the same way. Weights sum to 1.
PHASE_WEIGHTS: Final[dict[str, tuple[tuple[str, float], ...]]] = {
    "inventory_platform": (
        ("Diseño UI/UX", 0.23),
        ("Backend y API de inventario", 0.34),
        ("Autenticación y roles", 0.11),
        ("Cuadro de mando y reporting", 0.17),
        ("Pruebas y QA", 0.15),
    ),
    "mobile_ecommerce": (
        ("Diseño UI/UX móvil (iOS y Android)", 0.25),
        ("Catálogo de producto y buscador", 0.16),
        ("Carrito y proceso de compra", 0.14),
        ("Integración de pasarela de pago", 0.12),
        ("Gestión de pedidos", 0.11),
        ("Autenticación y perfil de usuario", 0.08),
        ("QA y publicación en stores", 0.14),
    ),
    "analytics_dashboard": (
        ("Diseño del cuadro de mando", 0.14),
        ("Desarrollo frontend", 0.22),
        ("Integración de datos y APIs", 0.16),
        ("Procesamiento en tiempo real", 0.19),
        ("Componentes de visualización", 0.12),
        ("Optimización de base de datos", 0.09),
        ("Pruebas y ajuste de rendimiento", 0.08),
    ),
    "crm": (
        ("Arquitectura y modelo de datos", 0.07),
        ("Diseño UI/UX", 0.12),
        ("Módulo de clientes", 0.15),
        ("Pipeline comercial y leads", 0.17),
        ("Herramientas de comunicación", 0.12),
        ("Informes y analítica", 0.10),
        ("Permisos y autenticación", 0.07),
        ("Integraciones con terceros", 0.09),
        ("Pruebas y QA", 0.11),
    ),
    "cms": (
        ("Arquitectura del CMS", 0.11),
        ("Panel de administración", 0.17),
        ("Editor de contenidos enriquecido", 0.16),
        ("Usuarios y permisos", 0.09),
        ("Biblioteca multimedia", 0.11),
        ("Programación de publicaciones", 0.08),
        ("Herramientas SEO", 0.06),
        ("Sistema de plantillas", 0.13),
        ("Pruebas y documentación", 0.09),
    ),
    "api_integration": (
        ("Arquitectura y planificación", 0.08),
        ("Integración de APIs de terceros", 0.25),
        ("Motor de sincronización de datos", 0.16),
        ("Gestor de webhooks", 0.11),
        ("Control de errores y reintentos", 0.09),
        ("Logging y monitorización", 0.08),
        ("Rate limiting", 0.06),
        ("Documentación técnica", 0.06),
        ("Pruebas de carga", 0.11),
    ),
}

#: Recommended team and duration, lifted from ESTIMATION_EXAMPLES so proposals
#: and transcripts quote the same shape of answer the model is primed to give.
PROJECT_TEAM: Final[dict[str, tuple[str, str]]] = {
    "inventory_platform": ("2 desarrolladores full-stack + 1 diseñador UX a tiempo parcial", "6-8 semanas"),
    "mobile_ecommerce": ("2 desarrolladores móviles (iOS y Android) + 1 backend + 1 QA", "10-12 semanas"),
    "analytics_dashboard": ("1 frontend + 1 backend + 1 ingeniero de datos", "8-10 semanas"),
    "crm": ("3 full-stack + 1 arquitecto de datos + 1 diseñador UI/UX + 1 QA", "12-14 semanas"),
    "cms": ("2 full-stack + 1 frontend + 1 diseñador UX + 1 QA", "10-12 semanas"),
    "api_integration": ("2 desarrolladores backend + 1 DevOps + 1 QA", "10-12 semanas"),
}


@dataclass(frozen=True)
class BudgetSpec:
    """One budget file before any defect is applied.

    ``hours`` drives everything monetary: the phases and ``total_amount`` are
    derived, never typed in, so a clean record is arithmetically consistent by
    construction and any inconsistency in the corpus is a planted one.
    """

    number: int  # -> budget_NNN.json
    budget_id: str
    client_code: str
    project_type: str
    hours: int
    signed_at: str  # ISO 8601; defects override it with other formats
    status: str  # draft | signed | rejected


#: 001-018 are canonical and clean (64% of the corpus). 019-028 carry the
#: planted defects; see PLANTED_DEFECTS for exactly which.
BUDGET_SPECS: Final[tuple[BudgetSpec, ...]] = (
    # -- clean majority ----------------------------------------------------
    BudgetSpec(1, "BUDGET-2023-0042", "CLI-1008", "analytics_dashboard", 255, "2023-09-18", "signed"),
    BudgetSpec(2, "BUDGET-2024-0117", "CLI-1002", "mobile_ecommerce", 320, "2024-01-23", "signed"),
    BudgetSpec(3, "BUDGET-2024-0205", "CLI-1001", "inventory_platform", 420, "2024-03-15", "signed"),
    BudgetSpec(4, "BUDGET-2024-0233", "CLI-1004", "crm", 410, "2024-07-09", "draft"),
    BudgetSpec(5, "BUDGET-2024-0241", "CLI-1006", "analytics_dashboard", 255, "2024-10-15", "signed"),
    BudgetSpec(6, "BUDGET-2024-0260", "CLI-1007", "cms", 300, "2024-04-11", "rejected"),
    BudgetSpec(7, "BUDGET-2025-0014", "CLI-1005", "cms", 320, "2025-02-18", "signed"),
    BudgetSpec(8, "BUDGET-2025-0031", "CLI-1010", "api_integration", 320, "2025-06-04", "draft"),
    BudgetSpec(9, "BUDGET-2025-0055", "CLI-1003", "analytics_dashboard", 180, "2025-03-07", "rejected"),
    BudgetSpec(10, "BUDGET-2025-0068", "CLI-1002", "crm", 360, "2025-04-22", "signed"),
    BudgetSpec(11, "BUDGET-2025-0082", "CLI-1009", "cms", 260, "2025-05-29", "draft"),
    BudgetSpec(12, "BUDGET-2025-0094", "CLI-1004", "mobile_ecommerce", 280, "2025-07-16", "rejected"),
    BudgetSpec(13, "BUDGET-2025-0103", "CLI-1008", "inventory_platform", 340, "2025-08-05", "signed"),
    BudgetSpec(14, "BUDGET-2025-0119", "CLI-1006", "api_integration", 220, "2025-09-10", "draft"),
    BudgetSpec(15, "BUDGET-2025-0126", "CLI-1007", "analytics_dashboard", 200, "2025-10-02", "draft"),
    BudgetSpec(16, "BUDGET-2025-0137", "CLI-1005", "inventory_platform", 290, "2025-11-13", "signed"),
    BudgetSpec(17, "BUDGET-2026-0008", "CLI-1010", "crm", 390, "2026-01-20", "signed"),
    BudgetSpec(18, "BUDGET-2026-0021", "CLI-1009", "mobile_ecommerce", 310, "2026-02-24", "draft"),
    # -- F1: same client, same date, same money, written four ways ---------
    BudgetSpec(19, "BUDGET-2024-0271", "CLI-1001", "inventory_platform", 1280, "2024-03-15", "signed"),
    BudgetSpec(20, "BUDGET-2024-0288", "CLI-1001", "crm", 1280, "2024-03-15", "signed"),
    BudgetSpec(21, "BUDGET-2024-0294", "CLI-1001", "analytics_dashboard", 900, "2024-06-18", "signed"),
    # -- F2: two pairs that collide on budget_id ---------------------------
    BudgetSpec(22, "BUDGET-2024-0312", "CLI-1003", "inventory_platform", 780, "2024-03-12", "signed"),
    BudgetSpec(23, "BUDGET-2024-0359", "CLI-1003", "inventory_platform", 840, "2024-05-02", "signed"),
    BudgetSpec(24, "BUDGET-2023-0188", "CLI-1009", "cms", 500, "2023-06-04", "signed"),
    BudgetSpec(25, "BUDGET-2023-0204", "CLI-1009", "cms", 478, "2023-09-18", "signed"),
    # -- F3 / F4 -----------------------------------------------------------
    BudgetSpec(26, "BUDGET-2025-0147", "CLI-1004", "crm", 300, "2025-01-15", "draft"),
    BudgetSpec(27, "BUDGET-2025-0158", "CLI-1007", "cms", 240, "2025-02-20", "draft"),
    BudgetSpec(28, "BUDGET-2024-0333", "CLI-1005", "api_integration", 260, "2024-11-28", "draft"),
)


# --------------------------------------------------------------------------
# The planted defects — ground truth for data/README.md and for later tests
# --------------------------------------------------------------------------

#: Marks "this row documents a defect but does not splice a value", used for
#: narrative defects inside transcripts. ``None`` cannot do that job, because
#: ``None`` is itself a value a defect might want to plant.
_NO_VALUE: Final[object] = object()


@dataclass(frozen=True)
class PlantedDefect:
    """A single documented defect.

    ``value`` doubles as the patch: when it is not ``_NO_VALUE`` the budget
    builder writes it straight into ``field``. That is what keeps the README
    table and the bytes on disk from ever disagreeing.
    """

    file: str  # path relative to data/corpus
    field: str
    family: str
    detail: str
    value: Any = _NO_VALUE


PLANTED_DEFECTS: Final[tuple[PlantedDefect, ...]] = (
    # ---- F1 Heterogeneidad de formato ------------------------------------
    PlantedDefect(
        "historical_budgets/budget_019.json", "client_name", F1,
        "Variante ortográfica de «ACME Corp.» (canónico en budget_003.json).",
        "Acme Corp",
    ),
    PlantedDefect(
        "historical_budgets/budget_019.json", "signed_at", F1,
        "Fecha en formato DD/MM/YYYY; misma fecha que budget_003.json («2024-03-15»).",
        "15/03/2024",
    ),
    PlantedDefect(
        "historical_budgets/budget_019.json", "currency", F1,
        "Moneda en minúsculas.",
        "eur",
    ),
    PlantedDefect(
        "historical_budgets/budget_019.json", "total_amount", F1,
        "Importe como string sin separadores (valor real 80000.0).",
        "80000",
    ),
    PlantedDefect(
        "historical_budgets/budget_020.json", "client_name", F1,
        "Variante ortográfica en minúsculas de «ACME Corp.».",
        "acme corp",
    ),
    PlantedDefect(
        "historical_budgets/budget_020.json", "client_code", F1,
        "Código de cliente en minúsculas (canónico «CLI-1001»).",
        "cli-1001",
    ),
    PlantedDefect(
        "historical_budgets/budget_020.json", "signed_at", F1,
        "Fecha en formato «Mon DD YYYY» en inglés; misma fecha que budget_003.json.",
        "Mar 15 2024",
    ),
    PlantedDefect(
        "historical_budgets/budget_020.json", "currency", F1,
        "Moneda como símbolo.",
        "€",
    ),
    PlantedDefect(
        "historical_budgets/budget_020.json", "total_amount", F1,
        "Importe como string con separador de miles y decimal europeos (valor real 80000.0).",
        "80.000,00",
    ),
    PlantedDefect(
        "historical_budgets/budget_021.json", "client_name", F1,
        "Variante abreviada de «ACME Corp.».",
        "ACME",
    ),
    PlantedDefect(
        "historical_budgets/budget_021.json", "currency", F1,
        "Moneda escrita en palabra.",
        "euros",
    ),
    PlantedDefect(
        "historical_budgets/budget_021.json", "total_amount", F1,
        "Importe como string sin separadores (valor real 56250.0).",
        "56250",
    ),
    PlantedDefect(
        "meeting_transcripts/2024-05-02_grupo-navarro_revision-presupuesto.txt", "(cuerpo del acta)", F1,
        "El importe aparece como «52.500,00 €», como «52500 EUR» y escrito en letra; "
        "la fecha aparece como «02/05/2024» y como «2 de mayo».",
    ),
    PlantedDefect(
        "meeting_transcripts/2023-02-27_grupo-navarro_seguimiento.txt", "(cuerpo del acta)", F1,
        "El mismo cliente se nombra «Grupo Navarro Logística», «grupo navarro» y «GN Logistica»; "
        "la moneda aparece como «euros» y como «EUR».",
    ),
    # ---- F2 Duplicados divergentes ---------------------------------------
    PlantedDefect(
        "historical_budgets/budget_023.json", "budget_id", F2,
        "Mismo budget_id que budget_022.json. Importes divergentes: 48750.0 (022) vs 52500.0 (023); "
        "signed_at 2024-03-12 (022) vs 2024-05-02 (023), así que «quedarse con el más reciente» "
        "resuelve a budget_023.json.",
        "BUDGET-2024-0312",
    ),
    PlantedDefect(
        "historical_budgets/budget_025.json", "budget_id", F2,
        "Mismo budget_id que budget_024.json. Importes divergentes: 31250.0 (024) vs 29875.0 (025); "
        "signed_at 2023-06-04 (024) vs 2023-09-18 (025), así que «quedarse con el más reciente» "
        "resuelve a budget_025.json.",
        "BUDGET-2023-0188",
    ),
    PlantedDefect(
        "meeting_transcripts/2024-05-02_grupo-navarro_revision-presupuesto.txt", "(cuerpo del acta)", F2,
        "El acta cita BUDGET-2024-0312 por 52.500,00 €, que contradice budget_022.json (48750.0) "
        "y coincide con budget_023.json (52500.0).",
    ),
    # ---- F3 Nulos disfrazados --------------------------------------------
    PlantedDefect(
        "historical_budgets/budget_026.json", "client_name", F3,
        "Centinela «N/A» en lugar del nombre del cliente (es CLI-1004, Farmacias Delgado).",
        "N/A",
    ),
    PlantedDefect(
        "historical_budgets/budget_026.json", "contact_email", F3,
        "Centinela «unknown» en lugar de un correo.",
        "unknown",
    ),
    PlantedDefect(
        "historical_budgets/budget_026.json", "account_manager", F3,
        "Centinela «TBD» en lugar del responsable de cuenta.",
        "TBD",
    ),
    PlantedDefect(
        "historical_budgets/budget_027.json", "client_name", F3,
        "Centinela «-» en lugar del nombre del cliente (es CLI-1007, Inmobiliaria Peñalver).",
        "-",
    ),
    PlantedDefect(
        "historical_budgets/budget_027.json", "contact_email", F3,
        "Cadena vacía.",
        "",
    ),
    PlantedDefect(
        "historical_budgets/budget_027.json", "account_manager", F3,
        "Centinela «pendiente» en lugar del responsable de cuenta.",
        "pendiente",
    ),
    PlantedDefect(
        "historical_budgets/budget_028.json", "client_name", F3,
        "Cadena de un solo espacio: parece rellena hasta que se hace strip() (es CLI-1005, Bodegas Villanueva).",
        " ",
    ),
    PlantedDefect(
        "meeting_transcripts/2023-02-27_grupo-navarro_seguimiento.txt", "(bloque de asistentes)", F3,
        "La cabecera del acta trae «Responsable de cuenta: pendiente», «Correo de contacto: N/A» "
        "y «Teléfono: -».",
    ),
    # ---- F4 Fuera de rango -----------------------------------------------
    PlantedDefect(
        "historical_budgets/budget_026.json", "total_amount", F4,
        "Importe negativo. Además ya no cuadra con phases[] ni con hours_estimated * 62.50 (= 18750.0).",
        -18750.0,
    ),
    PlantedDefect(
        "historical_budgets/budget_027.json", "total_amount", F4,
        "Importe absurdo (99 millones) frente a las 240 h del propio registro (= 15000.0).",
        99000000,
    ),
    PlantedDefect(
        "historical_budgets/budget_028.json", "hours_estimated", F4,
        "Horas en millones; phases[] sigue sumando 260 h.",
        1480000,
    ),
    PlantedDefect(
        "historical_budgets/budget_028.json", "signed_at", F4,
        "Fecha de firma en el futuro.",
        "2031-07-14",
    ),
    PlantedDefect(
        "historical_budgets/budget_028.json", "budget_id", F4,
        "budget_id fuera del patrón canónico BUDGET-YYYY-NNNN.",
        "BDG_2024_12",
    ),
)


# --------------------------------------------------------------------------
# Determinism helpers
# --------------------------------------------------------------------------

_PDF_ID_RE: Final[re.Pattern[bytes]] = re.compile(rb"/ID\[<([0-9A-Fa-f]+)><([0-9A-Fa-f]+)>\]")


def _freeze_pdf_id(raw: bytes, key: str) -> bytes:
    """Replace PyMuPDF's random trailer ``/ID`` with one derived from ``key``.

    PyMuPDF seeds the two file identifiers from the system RNG, which is the
    only non-deterministic part of its output. The replacement keeps the exact
    byte length of each identifier so every offset in the xref table — and
    ``startxref`` itself — stays valid without re-saving the document.
    """
    digest = hashlib.sha256(key.encode("utf-8")).hexdigest().upper().encode("ascii")

    def _replace(match: re.Match[bytes]) -> bytes:
        first, second = match.group(1), match.group(2)
        pool = digest * (1 + (len(first) + len(second)) // len(digest))
        return b"/ID[<" + pool[: len(first)] + b"><" + pool[: len(second)] + b">]"

    return _PDF_ID_RE.sub(_replace, raw, count=1)


def _freeze_core_xml(entry: bytes) -> bytes:
    """Pin the ``dcterms:created`` / ``dcterms:modified`` values in core.xml.

    openpyxl overwrites ``properties.modified`` with the current UTC time
    inside ``save()``, so setting it on the workbook is not enough: the stamp
    has to be removed after the library has written the bytes.
    """
    frozen = DOC_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    for tag in (b"created", b"modified"):
        entry = re.sub(
            rb"(<dcterms:" + tag + rb"[^>]*>)[^<]*(</dcterms:" + tag + rb">)",
            rb"\g<1>" + frozen + rb"\g<2>",
            entry,
        )
    return entry


def _freeze_zip(raw: bytes) -> bytes:
    """Rewrite an OOXML container with fixed per-entry timestamps.

    docx and xlsx are ZIPs, and ``ZipFile`` stamps each entry with the wall
    clock, so two identical documents written a second apart differ in bytes.
    Entry order and contents are preserved; only the timestamps are pinned.
    """
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                payload = _freeze_core_xml(payload)
            frozen = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            frozen.compress_type = zipfile.ZIP_DEFLATED
            frozen.external_attr = info.external_attr
            frozen.create_system = info.create_system
            dst.writestr(frozen, payload)
    return out.getvalue()


def _write(path: Path, payload: bytes, *, mtime: datetime = FRESH_MTIME) -> Path:
    """Write ``payload`` and pin the file's mtime, so re-runs are no-op diffs."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)
    stamp = mtime.timestamp()
    os.utime(path, (stamp, stamp))
    return path


def _eur(amount: float) -> str:
    """Format money the Spanish way (1.234,50) for the human-readable sources."""
    whole = f"{amount:,.2f}"
    return whole.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


# --------------------------------------------------------------------------
# Source 1 — historical budgets (JSON)
# --------------------------------------------------------------------------


def _phases_for(project_type: str, hours: int) -> list[dict[str, Any]]:
    """Split ``hours`` across the project's phases at the house rate.

    The last phase absorbs the rounding remainder so ``sum(phase hours)`` is
    exactly ``hours``; otherwise a clean record would fail its own arithmetic
    check for reasons that have nothing to do with the planted defects.
    """
    weights = PHASE_WEIGHTS[project_type]
    phases: list[dict[str, Any]] = []
    allocated = 0
    for index, (name, weight) in enumerate(weights):
        phase_hours = hours - allocated if index == len(weights) - 1 else round(hours * weight)
        allocated += phase_hours
        phases.append(
            {
                "name": name,
                "hours": phase_hours,
                "cost_eur": round(phase_hours * RATE_EUR_PER_HOUR, 2),
            }
        )
    return phases


def _budget_record(spec: BudgetSpec) -> dict[str, Any]:
    """Build the canonical record, then splice in whatever defects target it."""
    client = CLIENTS[spec.client_code]
    manager = ACCOUNT_MANAGERS[(spec.number - 1) % len(ACCOUNT_MANAGERS)]
    record: dict[str, Any] = {
        "budget_id": spec.budget_id,
        "client_name": client.name,
        "client_code": client.code,
        "project_type": spec.project_type,
        "total_amount": round(spec.hours * RATE_EUR_PER_HOUR, 2),
        "currency": "EUR",
        "hours_estimated": spec.hours,
        "signed_at": spec.signed_at,
        "status": spec.status,
        "contact_email": client.contact_email,
        "contact_phone": client.contact_phone,
        "account_manager": manager,
        "phases": _phases_for(spec.project_type, spec.hours),
    }

    filename = f"historical_budgets/budget_{spec.number:03d}.json"
    for defect in PLANTED_DEFECTS:
        if defect.file == filename and defect.value is not _NO_VALUE:
            record[defect.field] = defect.value
    return record


def build_budgets() -> list[Path]:
    """Write ``historical_budgets/budget_NNN.json``, one budget per file."""
    written: list[Path] = []
    for spec in BUDGET_SPECS:
        record = _budget_record(spec)
        payload = json.dumps(record, ensure_ascii=False, indent=2) + "\n"
        written.append(
            _write(
                CORPUS_DIR / "historical_budgets" / f"budget_{spec.number:03d}.json",
                payload.encode("utf-8"),
            )
        )
    return written


# --------------------------------------------------------------------------
# Source 2 — meeting transcripts (TXT, Spanish, two format eras)
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class MeetingSpec:
    """One meeting. ``date`` alone decides the format era (2024 = speaker tags)."""

    date: str  # ISO
    client_code: str
    project_type: str
    kind: str  # kickoff | alcance | revision | seguimiento
    budget_id: str
    hours: int


MEETING_SPECS: Final[tuple[MeetingSpec, ...]] = (
    # 2024+ : timestamped, one speaker per line
    MeetingSpec("2024-01-23", "CLI-1002", "mobile_ecommerce", "kickoff", "BUDGET-2024-0117", 320),
    MeetingSpec("2024-03-12", "CLI-1003", "inventory_platform", "kickoff", "BUDGET-2024-0312", 780),
    MeetingSpec("2024-05-02", "CLI-1003", "inventory_platform", "revision", "BUDGET-2024-0312", 840),
    MeetingSpec("2024-07-09", "CLI-1004", "crm", "alcance", "BUDGET-2024-0233", 410),
    MeetingSpec("2024-10-15", "CLI-1006", "analytics_dashboard", "kickoff", "BUDGET-2024-0241", 255),
    MeetingSpec("2025-02-18", "CLI-1005", "cms", "alcance", "BUDGET-2025-0014", 320),
    MeetingSpec("2025-06-04", "CLI-1010", "api_integration", "kickoff", "BUDGET-2025-0031", 320),
    # pre-2024 : free-running paragraphs, no speaker tags
    MeetingSpec("2022-04-19", "CLI-1001", "inventory_platform", "alcance", "BUDGET-2022-0077", 420),
    MeetingSpec("2022-11-08", "CLI-1007", "crm", "kickoff", "BUDGET-2022-0131", 380),
    MeetingSpec("2023-02-27", "CLI-1003", "analytics_dashboard", "seguimiento", "BUDGET-2023-0059", 500),
    MeetingSpec("2023-06-04", "CLI-1009", "cms", "alcance", "BUDGET-2023-0188", 500),
    MeetingSpec("2023-09-18", "CLI-1008", "analytics_dashboard", "kickoff", "BUDGET-2023-0042", 255),
)

MEETING_FILENAMES: Final[dict[str, str]] = {
    "2024-01-23": "2024-01-23_textiles-moreno_kickoff.txt",
    "2024-03-12": "2024-03-12_grupo-navarro_kickoff.txt",
    "2024-05-02": "2024-05-02_grupo-navarro_revision-presupuesto.txt",
    "2024-07-09": "2024-07-09_farmacias-delgado_alcance.txt",
    "2024-10-15": "2024-10-15_clinica-sanchis_kickoff.txt",
    "2025-02-18": "2025-02-18_bodegas-villanueva_alcance.txt",
    "2025-06-04": "2025-06-04_talleres-beltran_kickoff.txt",
    "2022-04-19": "2022-04-19_acme-corp_alcance.txt",
    "2022-11-08": "2022-11-08_inmobiliaria-penalver_kickoff.txt",
    "2023-02-27": "2023-02-27_grupo-navarro_seguimiento.txt",
    "2023-06-04": "2023-06-04_editorial-marimon_alcance.txt",
    "2023-09-18": "2023-09-18_distribuciones-arroyo_kickoff.txt",
}

#: What the client says is wrong today, per project type.
PROJECT_PAIN: Final[dict[str, str]] = {
    "inventory_platform": (
        "Ahora mismo el stock vive en tres hojas de cálculo distintas, una por almacén, "
        "y nadie sabe qué hay realmente hasta que alguien baja a contarlo."
    ),
    "mobile_ecommerce": (
        "Hoy vendemos por teléfono y por WhatsApp. Queremos una app propia con catálogo, "
        "carrito y pago, en iOS y en Android, sin depender de un marketplace."
    ),
    "analytics_dashboard": (
        "Los informes de ventas los montamos a mano todos los lunes y llegan tarde. "
        "Necesitamos un cuadro de mando que se actualice solo."
    ),
    "crm": (
        "Los contactos están repartidos entre Outlook y un Excel compartido. "
        "Perdemos oportunidades por pura desorganización, no por falta de clientes."
    ),
    "cms": (
        "Publicar una noticia hoy implica abrir un ticket a sistemas y esperar dos días. "
        "Queremos que marketing publique sin depender de nadie."
    ),
    "api_integration": (
        "Tenemos el ERP, la tienda online y el software de taller sin hablarse entre ellos. "
        "Todo se pasa a mano dos veces al día y se nos cuelan errores."
    ),
}

#: (technical question from the consultancy, client's answer).
PROJECT_TECH: Final[dict[str, tuple[str, str]]] = {
    "inventory_platform": (
        "¿Qué ERP tenéis detrás y cómo se consultan hoy las existencias?",
        "Un Sage antiguo, con acceso por ODBC. Las existencias se exportan a CSV cada noche.",
    ),
    "mobile_ecommerce": (
        "¿Con qué pasarela de pago trabajáis y qué catálogo tenéis ya digitalizado?",
        "Redsys para la tienda física y estamos abiertos a Stripe. El catálogo está en PrestaShop.",
    ),
    "analytics_dashboard": (
        "¿Dónde están los datos de ventas ahora mismo y con qué frecuencia cambian?",
        "En PostgreSQL y en un par de ficheros de Google Sheets. Cambian a cada minuto en campaña.",
    ),
    "crm": (
        "¿Qué volumen de contactos manejáis y qué necesitáis integrar con el correo?",
        "Unos 14.000 contactos y el correo corporativo es Microsoft 365, eso es innegociable.",
    ),
    "cms": (
        "¿Qué tipo de contenido publicáis y cuántas personas van a editar a la vez?",
        "Noticias, catálogo de títulos y bastante multimedia. Serían cinco editores y dos revisores.",
    ),
    "api_integration": (
        "¿Cuántos sistemas hay que sincronizar y cuáles exponen API hoy?",
        "Tres: ERP, tienda y taller. Solo la tienda tiene API REST; el resto es ficheros y SQL.",
    ),
}

#: Generic turns dropped in at seeded positions, tagged by who says them.
FILLER_TURNS: Final[tuple[tuple[str, str], ...]] = (
    ("client", "Una cosa importante: en campaña no podemos permitirnos caídas de servicio."),
    ("client", "¿El mantenimiento posterior entra en el precio o se contrata aparte?"),
    ("am", "El mantenimiento correctivo de los tres primeros meses va incluido; el evolutivo se contrata por bolsa de horas."),
    ("tech", "Lo montaríamos sobre Python y FastAPI en el backend, con React en el frontend."),
    ("tech", "Proponemos entornos separados de desarrollo, preproducción y producción desde el primer sprint."),
    ("client", "Necesitamos poder exportar todo a Excel, el comité financiero no va a cambiar de herramienta."),
    ("am", "Haremos una demo cada dos semanas para que no haya sorpresas al final."),
    ("tech", "El tratamiento de datos personales quedará conforme al RGPD, con registro de accesos."),
    ("client", "¿Quién formaría al equipo interno cuando esto esté en marcha?"),
    ("am", "Incluimos dos sesiones de formación y un manual de usuario en castellano."),
    ("tech", "Para la migración inicial necesitaremos un volcado de datos de prueba cuanto antes."),
    ("client", "Lo que no queremos es otro proyecto que se eternice. Eso ya nos pasó."),
)


def _meeting_turns(spec: MeetingSpec, rng: random.Random) -> list[tuple[str, str]]:
    """Build the turn list for a meeting as (speaker name, text) pairs.

    A fixed skeleton carries the facts a later extractor is supposed to find —
    budget id, hours, money, team, duration, PII — and seeded filler is woven
    around it so the twelve transcripts do not read like twelve copies.
    """
    client = CLIENTS[spec.client_code]
    label = PROJECT_LABELS[spec.project_type]
    manager = ACCOUNT_MANAGERS[rng.randrange(len(ACCOUNT_MANAGERS))]
    tech_lead = TECH_LEADS[rng.randrange(len(TECH_LEADS))]
    team, duration = PROJECT_TEAM[spec.project_type]
    question, answer = PROJECT_TECH[spec.project_type]
    total = spec.hours * RATE_EUR_PER_HOUR
    phases = ", ".join(name for name, _ in PHASE_WEIGHTS[spec.project_type][:4])

    speakers = {"am": manager, "tech": tech_lead, "client": client.contact_name, "client2": client.second_attendee}

    skeleton: list[tuple[str, str]] = [
        ("am", f"Buenos días y gracias por el tiempo. Soy {manager}, de {VENDOR_NAME}, "
               f"y me acompaña {tech_lead}, que llevará la parte técnica. El objetivo de hoy es cerrar "
               f"el alcance del proyecto de {label} para {client.name}."),
        ("client", f"Encantado. Yo soy {client.contact_name} y conmigo está {client.second_attendee}, "
                   f"de nuestro equipo de {rng.choice(('operaciones', 'sistemas', 'administración'))}."),
        ("client", PROJECT_PAIN[spec.project_type]),
        ("tech", question),
        ("client2", answer),
        ("am", f"Con eso, el alcance que planteamos se estructura en fases: {phases}, entre otras. "
               f"La referencia interna del presupuesto es {spec.budget_id}."),
        ("client", "¿Y de qué cifra estamos hablando?"),
        ("am", f"La estimación son {spec.hours} horas, que a nuestra tarifa de {_eur(RATE_EUR_PER_HOUR)} EUR/hora "
               f"salen {_eur(total)} EUR sin IVA."),
        ("client", f"¿{_eur(total)} euros? Tendré que pasarlo por dirección, pero no está lejos de lo que manejábamos."),
        ("tech", f"El equipo recomendado es {team} y la duración estimada {duration}."),
        ("client2", f"Nuestra fecha límite es {rng.choice(('la campaña de primavera', 'el cierre de ejercicio', 'la feria del sector', 'la vuelta de verano'))}. "
                    f"Si se pasa de ahí, pierde sentido."),
        ("am", f"Anotado. Os envío el acta y la propuesta formal a {client.contact_email} esta misma tarde."),
        ("client", f"Perfecto. Si hay cualquier urgencia, mi móvil directo es {client.contact_phone}."),
        ("am", f"Y el nuestro, {VENDOR_PHONE}, o escribid a {VENDOR_EMAIL}. Cerramos aquí."),
    ]

    filler = rng.sample(FILLER_TURNS, k=rng.randint(4, 6))
    turns = list(skeleton)
    # Insert filler between the scope discussion and the closing, so the facts
    # in the skeleton keep their relative order.
    for role, text in filler:
        position = rng.randint(4, len(turns) - 3)
        turns.insert(position, (role, text))

    # Splice the planted turns in before the two closing turns, so the defects
    # sit in the body of the meeting rather than after the goodbyes.
    if spec.kind == "revision":
        turns[-2:-2] = _revision_defect_turns()
    if spec.kind == "seguimiento":
        turns[-2:-2] = _seguimiento_defect_turns()

    return [(speakers[role], text) for role, text in turns]


def _revision_defect_turns() -> list[tuple[str, str]]:
    """The F1 + F2 turns of the 2024-05-02 budget-revision meeting.

    The same figure is spoken three ways and contradicts budget_022.json, which
    is what gives a later reconciliation step something real to resolve.
    """
    return [
        ("am", "Recapitulando el presupuesto BUDGET-2024-0312: la versión que firmasteis el 12/03/2024 "
               "iba por 48.750,00 € y la revisión de hoy, 02/05/2024, queda en 52.500,00 €."),
        ("client", "Es decir, cincuenta y dos mil quinientos euros. Que conste que yo tenía apuntado 52500 EUR "
                   "desde la reunión del 2 de mayo."),
        ("am", "Correcto. Mantenemos el mismo identificador de presupuesto, BUDGET-2024-0312, para no abrir otro expediente."),
    ]


def _seguimiento_defect_turns() -> list[tuple[str, str]]:
    """The F1 + F3 turns of the 2023-02-27 follow-up meeting."""
    return [
        ("am", "Para el acta: el cliente figura como Grupo Navarro Logística en contrato, como grupo navarro "
               "en la herramienta de tickets y como GN Logistica en el ERP. Habría que unificarlo."),
        ("client", "Y el importe que yo tengo anotado son treinta y un mil doscientos cincuenta euros, "
                   "aunque en el PDF pone 31.250,00 EUR."),
    ]


def _render_timestamped(turns: list[tuple[str, str]], spec: MeetingSpec, rng: random.Random) -> str:
    """2024-era format: ``[hh:mm:ss] Nombre Apellido: texto``, one turn per line."""
    client = CLIENTS[spec.client_code]
    seconds = 10 * 3600  # meetings start at 10:00:00
    lines = [
        f"# Acta de reunión — {client.name} — {spec.date}",
        f"# Transcripción automática · {VENDOR_NAME}",
        "",
    ]
    for speaker, text in turns:
        stamp = f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"
        lines.append(f"[{stamp}] {speaker}: {text}")
        seconds += rng.randint(25, 95)
    lines.append("")
    return "\n".join(lines)


def _render_freeform(turns: list[tuple[str, str]], spec: MeetingSpec, rng: random.Random) -> str:
    """Pre-2024 format: prose, no speaker tags, inconsistent whitespace.

    The irregularity is the point — these files are what a chunker trips over,
    so the blank-line runs and double spaces are seeded rather than tidy.
    """
    client = CLIENTS[spec.client_code]
    header = [
        f"ACTA DE REUNION - {client.name.upper()} - {spec.date}",
        f"{VENDOR_NAME}   Asistentes y notas tomadas a mano, sin grabacion.",
    ]
    if spec.kind == "seguimiento":
        header += [
            "Responsable de cuenta: pendiente",
            "Correo de contacto: N/A",
            "Teléfono: -",
        ]

    body: list[str] = []
    paragraph: list[str] = []
    verbs = ("comenta que", "explica que", "añade que", "insiste en que", "apunta que", "recuerda que")
    for index, (speaker, text) in enumerate(turns):
        first_name = speaker.split()[0]
        head = 1 if text[0] in "¿¡" else 0
        quoted = text[:head] + text[head].lower() + text[head + 1 :]
        sentence = f"{first_name} {rng.choice(verbs)} {quoted}"
        if rng.random() < 0.35:
            # Sporadic double spacing, exactly the kind of thing a naive
            # whitespace split mishandles.
            sentence = sentence.replace(". ", ".  ", 1)
        paragraph.append(sentence)
        if len(paragraph) >= rng.randint(2, 4) or index == len(turns) - 1:
            body.append(" ".join(paragraph))
            body.append("\n" * rng.randint(0, 2))
            paragraph = []

    tail = "Pendiente de confirmar presupuesto y calendario.   Se enviara propuesta por correo."
    return "\n\n".join(header) + "\n\n" + "\n".join(body) + "\n" + tail + "\n"


def build_transcripts() -> list[Path]:
    """Write the twelve transcripts, split across the two format eras."""
    written: list[Path] = []
    for index, spec in enumerate(MEETING_SPECS):
        rng = random.Random(SEED + 1000 + index)
        turns = _meeting_turns(spec, rng)
        era_2024 = spec.date >= "2024-01-01"
        text = (
            _render_timestamped(turns, spec, rng)
            if era_2024
            else _render_freeform(turns, spec, rng)
        )
        written.append(
            _write(
                CORPUS_DIR / "meeting_transcripts" / MEETING_FILENAMES[spec.date],
                text.encode("utf-8"),
            )
        )
    return written


# --------------------------------------------------------------------------
# Source 3 — proposal templates (DOCX)
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class ProposalSpec:
    """One Word proposal. ``budget_id`` ties it back to a clean budget file."""

    filename: str
    client_code: str
    project_type: str
    budget_id: str
    hours: int


PROPOSAL_SPECS: Final[tuple[ProposalSpec, ...]] = (
    ProposalSpec("propuesta_inventario_acme.docx", "CLI-1001", "inventory_platform", "BUDGET-2024-0205", 420),
    ProposalSpec("propuesta_ecommerce_textiles-moreno.docx", "CLI-1002", "mobile_ecommerce", "BUDGET-2024-0117", 320),
    ProposalSpec("propuesta_dashboard_clinica-sanchis.docx", "CLI-1006", "analytics_dashboard", "BUDGET-2024-0241", 255),
    ProposalSpec("propuesta_crm_textiles-moreno.docx", "CLI-1002", "crm", "BUDGET-2025-0068", 360),
    ProposalSpec("propuesta_cms_bodegas-villanueva.docx", "CLI-1005", "cms", "BUDGET-2025-0014", 320),
)


def _build_proposal(spec: ProposalSpec) -> bytes:
    """Render one proposal with a real heading hierarchy.

    Headings are ``add_heading`` levels rather than bold paragraphs because the
    downstream splitter keys on the Heading styles; a visually identical
    document made of bold runs would silently produce one giant section.
    """
    client = CLIENTS[spec.client_code]
    label = PROJECT_LABELS[spec.project_type]
    team, duration = PROJECT_TEAM[spec.project_type]
    phases = _phases_for(spec.project_type, spec.hours)
    total = sum(phase["cost_eur"] for phase in phases)

    document = Document()
    document.add_heading(f"Propuesta de proyecto — {label.capitalize()}", level=0)
    subtitle = document.add_paragraph(
        f"{client.name} · {client.city}\nReferencia de presupuesto: {spec.budget_id}\n"
        f"Preparado por {VENDOR_NAME}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    document.add_heading("Alcance", level=1)
    document.add_paragraph(
        f"El presente documento describe el alcance del proyecto de {label} solicitado por {client.name}. "
        "El alcance se ha cerrado a partir de las reuniones de trabajo mantenidas con el equipo "
        "del cliente y sustituye a cualquier estimación previa."
    )
    document.add_heading("Dentro del alcance", level=2)
    for phase in phases:
        document.add_paragraph(phase["name"], style="List Bullet")
    document.add_heading("Fuera del alcance", level=2)
    for item in (
        "Migración de datos históricos anteriores a tres años.",
        "Licencias de software de terceros y costes de infraestructura cloud.",
        "Soporte en horario 24x7 durante el primer año.",
    ):
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Entregables", level=1)
    for item in (
        "Código fuente en el repositorio del cliente, con historial completo.",
        "Documentación técnica de despliegue y manual de usuario en castellano.",
        "Plan de pruebas y resultados de la batería de aceptación.",
        "Dos sesiones de formación para el equipo interno.",
    ):
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Cronograma", level=1)
    document.add_paragraph(
        f"La duración estimada es de {duration}, con demos quincenales y una entrega "
        "en preproducción antes de la aceptación final."
    )
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, ("Fase", "Horas", "Importe (EUR)")):
        cell.text = heading
        cell.paragraphs[0].runs[0].font.bold = True
    for phase in phases:
        row = table.add_row().cells
        row[0].text = phase["name"]
        row[1].text = str(phase["hours"])
        row[2].text = _eur(phase["cost_eur"])
    totals = table.add_row().cells
    totals[0].text = "TOTAL"
    totals[1].text = str(sum(phase["hours"] for phase in phases))
    totals[2].text = _eur(total)
    for cell in totals:
        cell.paragraphs[0].runs[0].font.bold = True

    document.add_heading("Equipo", level=1)
    document.add_paragraph(f"Equipo recomendado: {team}.")
    document.add_paragraph(
        "El equipo se asigna en exclusiva durante las fases de construcción y comparte "
        "disponibilidad con otros proyectos únicamente en la fase de garantía."
    )

    document.add_heading("Condiciones económicas", level=1)
    document.add_paragraph(
        f"Importe total: {_eur(total)} EUR (IVA no incluido), correspondientes a "
        f"{sum(phase['hours'] for phase in phases)} horas a una tarifa de "
        f"{_eur(RATE_EUR_PER_HOUR)} EUR/hora."
    )
    document.add_heading("Forma de pago", level=2)
    for item in (
        "30 % a la firma de la propuesta.",
        "40 % a la entrega en preproducción.",
        "30 % a la aceptación final.",
    ):
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Validez de la oferta", level=2)
    document.add_paragraph(
        "La presente oferta tiene una validez de 30 días naturales desde su emisión. "
        f"Para cualquier aclaración: {VENDOR_EMAIL} · {VENDOR_PHONE}."
    )

    for style_name in ("Normal",):
        document.styles[style_name].font.size = Pt(10)

    properties = document.core_properties
    properties.author = VENDOR_NAME
    properties.last_modified_by = VENDOR_NAME
    properties.created = DOC_TIMESTAMP
    properties.modified = DOC_TIMESTAMP
    properties.title = f"Propuesta {spec.budget_id}"
    properties.revision = 1

    buffer = io.BytesIO()
    document.save(buffer)
    return _freeze_zip(buffer.getvalue())


def build_proposals() -> list[Path]:
    """Write the five Word proposal templates."""
    return [
        _write(CORPUS_DIR / "proposal_templates" / spec.filename, _build_proposal(spec))
        for spec in PROPOSAL_SPECS
    ]


# --------------------------------------------------------------------------
# Source 4 — signed contracts (PDF with a real text layer)
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class ContractSpec:
    """One signed contract, pinned to a signed budget file.

    ``optional_clauses`` is how many of the six non-essential clauses this
    contract carries. It is a field rather than a random draw so the page count
    of each PDF is a stated property of the corpus, not an emergent one: the
    documents have to span 1-3 pages and a seeded draw could silently stop
    doing that after an unrelated edit to the prose.
    """

    filename: str
    client_code: str
    project_type: str
    budget_id: str
    hours: int
    signed_at: str
    optional_clauses: int


CONTRACT_SPECS: Final[tuple[ContractSpec, ...]] = (
    ContractSpec("contrato_BUDGET-2023-0042_distribuciones-arroyo.pdf", "CLI-1008", "analytics_dashboard", "BUDGET-2023-0042", 255, "18/09/2023", 6),
    ContractSpec("contrato_BUDGET-2024-0117_textiles-moreno.pdf", "CLI-1002", "mobile_ecommerce", "BUDGET-2024-0117", 320, "23/01/2024", 3),
    ContractSpec("contrato_BUDGET-2024-0205_acme-corp.pdf", "CLI-1001", "inventory_platform", "BUDGET-2024-0205", 420, "15/03/2024", 6),
    ContractSpec("contrato_BUDGET-2024-0241_clinica-sanchis.pdf", "CLI-1006", "analytics_dashboard", "BUDGET-2024-0241", 255, "15/10/2024", 0),
    ContractSpec("contrato_BUDGET-2025-0068_textiles-moreno.pdf", "CLI-1002", "crm", "BUDGET-2025-0068", 360, "22/04/2025", 4),
    ContractSpec("contrato_BUDGET-2025-0103_distribuciones-arroyo.pdf", "CLI-1008", "inventory_platform", "BUDGET-2025-0103", 340, "05/08/2025", 2),
)


class _PdfFlow:
    """Minimal flowing-text layout on top of ``insert_textbox``.

    PyMuPDF has no concept of a document flow: ``insert_textbox`` fills one
    rectangle and tells you only whether the text fitted. This measures each
    block with ``get_text_length`` to decide how tall it will be, and starts a
    new page when the cursor would run past the margin — enough structure for a
    contract, with every glyph in a real text layer so no OCR is ever needed.
    """

    LEFT, RIGHT, TOP, BOTTOM = 62.0, 533.0, 62.0, 780.0

    def __init__(self) -> None:
        self.doc = pymupdf.open()
        self.page = self.doc.new_page()
        self.cursor = self.TOP

    def _width(self) -> float:
        return self.RIGHT - self.LEFT

    def block(self, text: str, *, size: float = 10.0, bold: bool = False, space_after: float = 8.0) -> None:
        fontname = "hebo" if bold else "helv"
        leading = size * 1.35
        text_width = pymupdf.get_text_length(text, fontname=fontname, fontsize=size)
        # +1 line of slack: get_text_length measures an unwrapped string, and
        # word wrapping can push the last word onto an extra line.
        lines = int(text_width // self._width()) + 2
        height = lines * leading

        if self.cursor + height > self.BOTTOM:
            self.page = self.doc.new_page()
            self.cursor = self.TOP

        rect = pymupdf.Rect(self.LEFT, self.cursor, self.RIGHT, min(self.cursor + height + leading, self.BOTTOM))
        self.page.insert_textbox(rect, text, fontsize=size, fontname=fontname, lineheight=1.35)
        self.cursor += height + space_after

    def gap(self, amount: float = 12.0) -> None:
        self.cursor += amount

    def tobytes(self) -> bytes:
        data: bytes = self.doc.tobytes()
        self.doc.close()
        return data


def _build_contract(spec: ContractSpec) -> bytes:
    """Render one Spanish services contract as a born-digital PDF."""
    client = CLIENTS[spec.client_code]
    label = PROJECT_LABELS[spec.project_type]
    _, duration = PROJECT_TEAM[spec.project_type]
    total = spec.hours * RATE_EUR_PER_HOUR

    flow = _PdfFlow()
    flow.block("CONTRATO DE PRESTACIÓN DE SERVICIOS DE DESARROLLO DE SOFTWARE", size=12.5, bold=True, space_after=14)
    flow.block(f"Referencia de presupuesto: {spec.budget_id}", size=10, bold=True)
    flow.block(f"En Madrid, a {spec.signed_at}.", size=10, space_after=14)

    flow.block("REUNIDOS", size=11, bold=True)
    flow.block(
        f"De una parte, {VENDOR_NAME}, con CIF {VENDOR_CIF} y domicilio social en {VENDOR_ADDRESS}, "
        "en adelante EL PROVEEDOR, representada en este acto por su director de proyectos.",
    )
    flow.block(
        f"De otra parte, {client.name}, con domicilio en {client.city}, en adelante EL CLIENTE, "
        f"representada por D./Dña. {client.contact_name}, con correo electrónico {client.contact_email} "
        f"y teléfono {client.contact_phone}.",
    )
    flow.block(
        "Ambas partes se reconocen capacidad legal suficiente para contratar y obligarse, y a tal efecto",
        space_after=12,
    )

    flow.block("EXPONEN", size=11, bold=True)
    flow.block(
        f"I. Que EL CLIENTE precisa la construcción de un proyecto de {label} y ha solicitado oferta a EL PROVEEDOR."
    )
    flow.block(
        f"II. Que EL PROVEEDOR ha emitido el presupuesto {spec.budget_id}, aceptado por EL CLIENTE, "
        "cuyo contenido se incorpora como Anexo I al presente contrato."
    )
    flow.block("III. Que ambas partes acuerdan suscribir el presente contrato con arreglo a las siguientes", space_after=12)

    flow.block("CLÁUSULAS", size=11, bold=True)

    # The first four clauses are what makes the document a contract at all and
    # are always present; the rest are trimmed per spec.
    mandatory: tuple[tuple[str, str], ...] = (
        (
            "PRIMERA. Objeto",
            f"El objeto del contrato es el diseño, desarrollo, pruebas y puesta en producción del proyecto de {label} "
            f"para EL CLIENTE, conforme al alcance detallado en el presupuesto {spec.budget_id}. Queda excluida "
            "cualquier funcionalidad no recogida expresamente en dicho documento, que requerirá adenda escrita.",
        ),
        (
            "SEGUNDA. Precio y forma de pago",
            f"El precio total asciende a {_eur(total)} EUR, IVA no incluido, correspondientes a {spec.hours} horas "
            f"de trabajo a una tarifa unitaria de {_eur(RATE_EUR_PER_HOUR)} EUR/hora. El pago se realizará en tres "
            "hitos: un 30 % a la firma, un 40 % a la entrega en preproducción y un 30 % a la aceptación final. "
            "Las facturas se emitirán a 30 días fecha factura mediante transferencia bancaria.",
        ),
        (
            "TERCERA. Plazo de ejecución",
            f"El plazo estimado de ejecución es de {duration} desde la firma del presente contrato y la entrega "
            "del entorno y los accesos necesarios por parte de EL CLIENTE. Los retrasos imputables a EL CLIENTE "
            "en la validación de entregables prorrogarán el plazo en el mismo número de días naturales.",
        ),
        (
            "CUARTA. Equipo y seguimiento",
            "EL PROVEEDOR asignará el equipo recogido en el presupuesto y designará un responsable de cuenta único. "
            "Se celebrará una reunión de seguimiento quincenal, de la que se levantará acta remitida por correo "
            "electrónico a los interlocutores designados por ambas partes.",
        ),
    )
    optional: tuple[tuple[str, str], ...] = (
        (
            "QUINTA. Confidencialidad",
            "Ambas partes se obligan a mantener la más estricta confidencialidad sobre la información a la que "
            "accedan con ocasión del presente contrato, obligación que subsistirá durante cinco años desde su "
            "extinción, cualquiera que sea la causa.",
        ),
        (
            "SEXTA. Propiedad intelectual",
            "La titularidad del código fuente desarrollado a medida corresponderá a EL CLIENTE una vez satisfecho "
            "el precio íntegro. EL PROVEEDOR conservará la titularidad de sus componentes y librerías preexistentes, "
            "sobre los que concede a EL CLIENTE una licencia de uso indefinida, no exclusiva e intransferible.",
        ),
        (
            "SÉPTIMA. Protección de datos",
            "En la medida en que EL PROVEEDOR acceda a datos personales responsabilidad de EL CLIENTE, actuará como "
            "encargado del tratamiento conforme al Reglamento (UE) 2016/679, tratándolos únicamente siguiendo "
            "instrucciones documentadas de EL CLIENTE y suprimiéndolos a la finalización del servicio.",
        ),
        (
            "OCTAVA. Garantía",
            "EL PROVEEDOR garantiza la corrección sin coste de los defectos de software detectados durante los tres "
            "meses siguientes a la aceptación final, siempre que no deriven de modificaciones realizadas por terceros "
            "ni de un uso distinto al previsto.",
        ),
        (
            "NOVENA. Resolución",
            "Serán causas de resolución el incumplimiento grave de las obligaciones asumidas, el impago de dos "
            "facturas consecutivas y la declaración de concurso de cualquiera de las partes.",
        ),
        (
            "DÉCIMA. Legislación y jurisdicción",
            "El presente contrato se rige por la legislación española. Para cuantas cuestiones se susciten en su "
            "interpretación o cumplimiento, las partes se someten a los Juzgados y Tribunales de Madrid, con "
            "renuncia expresa a cualquier otro fuero que pudiera corresponderles.",
        ),
    )
    for title, body in mandatory + optional[: spec.optional_clauses]:
        flow.block(title, size=10, bold=True, space_after=3)
        flow.block(body, size=9.5, space_after=10)

    flow.gap(18)
    flow.block(
        "Y en prueba de conformidad, ambas partes firman el presente contrato por duplicado y a un solo efecto, "
        "en el lugar y fecha indicados en el encabezamiento.",
        space_after=24,
    )
    flow.block("Por EL PROVEEDOR                                        Por EL CLIENTE", size=10, bold=True, space_after=6)
    flow.block(f"{VENDOR_NAME}                      {client.name}", size=9.5, space_after=4)
    flow.block(f"{VENDOR_EMAIL}                                  {client.contact_email}", size=9.5)

    raw = flow.tobytes()
    return _freeze_pdf_id(raw, spec.filename)


def build_contracts() -> list[Path]:
    """Write the six signed contracts as born-digital PDFs."""
    return [
        _write(CORPUS_DIR / "signed_contracts" / spec.filename, _build_contract(spec))
        for spec in CONTRACT_SPECS
    ]


# --------------------------------------------------------------------------
# Source 5 — rate card (XLSX, deliberately stale)
# --------------------------------------------------------------------------

#: Role -> per-seniority rate. Centred on the house rate: the mean of the whole
#: grid is exactly 62.50 EUR/h, with seniority fanning out +-20 EUR around it.
#: ``main()`` asserts the mean, so a future edit cannot quietly drift off the
#: rate the few-shot examples assume.
RATE_CARD_ROWS: Final[tuple[tuple[str, float, float, float, float], ...]] = (
    ("Desarrollador Full-Stack", 40.00, 57.50, 70.00, 82.50),
    ("Desarrollador Frontend", 37.50, 55.00, 67.50, 80.00),
    ("Desarrollador Backend", 40.00, 57.50, 70.00, 82.50),
    ("Desarrollador Móvil", 42.50, 60.00, 72.50, 85.00),
    ("Ingeniero de Datos", 42.50, 60.00, 72.50, 85.00),
    ("Diseñador UX/UI", 35.00, 52.50, 65.00, 77.50),
    ("QA Engineer", 32.50, 50.00, 62.50, 75.00),
    ("DevOps / SRE", 42.50, 60.00, 72.50, 85.00),
    ("Arquitecto de Software", 47.50, 65.00, 77.50, 90.00),
    ("Jefe de Proyecto", 40.00, 57.50, 70.00, 82.50),
)


def build_rate_card() -> Path:
    """Write the single rate card and back-date it to January 2024.

    The mtime is the whole point of this file: it is the source the pipeline
    must *drop* for being out of date, and a staleness rule reads the
    filesystem, not the cell that says "Vigencia 2024".
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tarifas 2024"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDE5EE")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = "#,##0.00 \"€/h\""

    sheet["A1"] = f"{VENDOR_NAME} — Tarifario interno"
    sheet["A1"].font = Font(bold=True, size=13)
    sheet["A2"] = "Vigencia: 01/01/2024 – 31/12/2024"
    sheet["A3"] = f"Tarifa base de referencia: {RATE_EUR_PER_HOUR:.2f} €/h (todas las estimaciones parten de aquí)"
    sheet["A4"] = "Documento interno. Revisar antes de cada cierre de ejercicio."

    headers = ("Rol", "Junior", "Mid", "Senior", "Lead", "Media")
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=6, column=column, value=title)
        cell.font = bold
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for offset, (role, junior, mid, senior, lead) in enumerate(RATE_CARD_ROWS):
        row = 7 + offset
        sheet.cell(row=row, column=1, value=role).border = border
        for column, rate in enumerate((junior, mid, senior, lead), start=2):
            cell = sheet.cell(row=row, column=column, value=rate)
            cell.number_format = money
            cell.border = border
        average = sheet.cell(row=row, column=6, value=round((junior + mid + senior + lead) / 4, 2))
        average.number_format = money
        average.border = border

    last_row = 6 + len(RATE_CARD_ROWS)
    overall = round(
        sum(sum(rates) for _, *rates in RATE_CARD_ROWS) / (len(RATE_CARD_ROWS) * 4), 2
    )
    sheet.cell(row=last_row + 2, column=1, value="Tarifa media de la casa").font = bold
    mean_cell = sheet.cell(row=last_row + 2, column=6, value=overall)
    mean_cell.number_format = money
    mean_cell.font = bold
    sheet.cell(
        row=last_row + 3,
        column=1,
        value=f"Tarifa aplicada por defecto en presupuestos: {RATE_EUR_PER_HOUR:.2f} €/h",
    )

    sheet.column_dimensions["A"].width = 30
    for column in range(2, 7):
        sheet.column_dimensions[get_column_letter(column)].width = 12
    sheet.freeze_panes = "A7"

    workbook.properties.creator = VENDOR_NAME
    workbook.properties.lastModifiedBy = VENDOR_NAME
    workbook.properties.created = DOC_TIMESTAMP
    workbook.properties.modified = DOC_TIMESTAMP
    workbook.properties.title = "Tarifario 2024"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return _write(
        CORPUS_DIR / "rate_card" / "rate_card_2024.xlsx",
        _freeze_zip(buffer.getvalue()),
        mtime=STALE_MTIME,
    )


# --------------------------------------------------------------------------
# data/README.md — rendered from the same tuple that plants the defects
# --------------------------------------------------------------------------


def _describe_value(defect: PlantedDefect) -> str:
    if defect.value is _NO_VALUE:
        return "—"
    if isinstance(defect.value, str):
        return f"`{defect.value!r}`"
    return f"`{defect.value}`"


def _source_stats() -> list[tuple[str, int, int, str]]:
    """(relative dir, file count, total bytes, glob) for the README tree."""
    stats: list[tuple[str, int, int, str]] = []
    for folder, pattern in (
        ("historical_budgets", "*.json"),
        ("meeting_transcripts", "*.txt"),
        ("proposal_templates", "*.docx"),
        ("signed_contracts", "*.pdf"),
        ("rate_card", "*.xlsx"),
    ):
        files = sorted((CORPUS_DIR / folder).glob(pattern))
        stats.append((folder, len(files), sum(f.stat().st_size for f in files), pattern))
    return stats


def _clean_budget_files() -> list[str]:
    """Budget files carrying no defect at all, counted the strict way.

    A file is clean only if nothing was planted in it *and* its ``budget_id``
    is unique. The two uncorrupted halves of the duplicate pairs are therefore
    excluded: they are individually valid but they participate in a defect, so
    counting them as clean would overstate the majority a validation step is
    supposed to let through.
    """
    planted = {d.file for d in PLANTED_DEFECTS if d.file.startswith("historical_budgets/")}
    seen: dict[str, list[str]] = {}
    for spec in BUDGET_SPECS:
        name = f"historical_budgets/budget_{spec.number:03d}.json"
        seen.setdefault(_budget_record(spec)["budget_id"], []).append(name)
    colliding = {name for names in seen.values() if len(names) > 1 for name in names}
    return sorted(
        f"historical_budgets/budget_{spec.number:03d}.json"
        for spec in BUDGET_SPECS
        if f"historical_budgets/budget_{spec.number:03d}.json" not in planted | colliding
    )


def write_readme() -> Path:
    """Write ``data/README.md``, including the ground-truth defect table."""
    stats = _source_stats()
    total_bytes = sum(size for _, _, size, _ in stats)

    tree_lines = [
        "```",
        "data/",
        "├── build_corpus.py",
        "├── README.md",
        "└── corpus/",
    ]
    for index, (folder, count, size, pattern) in enumerate(stats):
        branch = "    └──" if index == len(stats) - 1 else "    ├──"
        tree_lines.append(f"{branch} {folder}/{' ' * max(1, 22 - len(folder))}{count:>2} × {pattern:<7} {size / 1024:7.1f} KB")
    tree_lines.append("```")

    family_counts: dict[str, int] = {}
    for defect in PLANTED_DEFECTS:
        family_counts[defect.family] = family_counts.get(defect.family, 0) + 1

    table_rows = [
        f"| `{defect.file}` | `{defect.field}` | {defect.family} | {_describe_value(defect)} | {defect.detail} |"
        for defect in PLANTED_DEFECTS
    ]

    clean = len(_clean_budget_files())

    content = f"""# `data/` — synthetic enterprise corpus

Everything under `data/corpus/` is **generated**, fake and safe to commit. No
real client, person, email, phone number or amount appears here.

The corpus exists so the ingestion/cleaning work can be tested against data
that is deliberately as messy as the real thing, with the mess written down.
It is coherent with the world of `app/context/examples.py`: a Spanish software
consultancy, EUR, hours billed at **{RATE_EUR_PER_HOUR:.2f} EUR/hour**, phases and
tasks, teams and durations in weeks.

## Regenerate

```bash
uv run python data/build_corpus.py
```

The generator is deterministic: fixed seed (`SEED = {SEED}`), fixed document
timestamps, no `datetime.now()` anywhere in the content. Two consecutive runs
produce **byte-identical files**, so re-running it on a clean checkout leaves an
empty diff. `data/corpus/` is wiped and rebuilt on every run, so it is also
idempotent in the stronger sense: nothing stale survives.

Two extra steps are needed for that guarantee and are easy to break by accident:
PyMuPDF stamps a random `/ID` into each PDF trailer, and ZIP containers (docx,
xlsx) record the wall clock per entry. Both are frozen after the library writes
the bytes.

## Contents

{chr(10).join(tree_lines)}

Total: **{total_bytes / 1024:.1f} KB**.

| Source | What it is |
|---|---|
| `historical_budgets/` | {stats[0][1]} JSON files, one budget per file, as exported from the ERP. Fields: `budget_id`, `client_name`, `client_code`, `project_type`, `total_amount`, `currency`, `hours_estimated`, `signed_at`, `status`, `contact_email`, `contact_phone`, `account_manager`, `phases[]` (`name`, `hours`, `cost_eur`). Canonical forms: `BUDGET-YYYY-NNNN`, `CLI-NNNN`, status ∈ {{draft, signed, rejected}}. |
| `meeting_transcripts/` | {stats[1][1]} Spanish transcripts of client scoping/kickoff meetings, in **two format eras**. 2024 and later ({len([s for s in MEETING_SPECS if s.date >= "2024-01-01"])} files): every line is `[hh:mm:ss] Nombre Apellido: texto`. Before 2024 ({len([s for s in MEETING_SPECS if s.date < "2024-01-01"])} files): no speaker tags, free-running paragraphs, irregular blank lines and double spaces. All of them carry inline PII — Spanish names, organisations, emails, phone numbers — which is what an anonymisation step has to catch. |
| `proposal_templates/` | {stats[2][1]} Spanish Word proposals with a real `add_heading` hierarchy (level 0 title, level 1 sections `Alcance` / `Entregables` / `Cronograma` / `Equipo` / `Condiciones económicas`, level 2 subsections) and a real table of phases, hours and amounts. A parser that splits by heading style will find them. |
| `signed_contracts/` | {stats[3][1]} Spanish services contracts as **born-digital** PDFs, 1–3 pages, with a genuine text layer. No scans, no image-only pages, no OCR required. Each references a `budget_id` and a client. |
| `rate_card/` | 1 XLSX, role × seniority, rates centred on {RATE_EUR_PER_HOUR:.2f} EUR/h. **Deliberately stale**: its mtime is set to {STALE_MTIME:%Y-%m-%d} while every other corpus file is stamped {FRESH_MTIME:%Y-%m-%d}, so a staleness rule reading the filesystem will exclude it. |

## Planted defects — ground truth

{clean} of the {len(BUDGET_SPECS)} budget files ({clean * 100 // len(BUDGET_SPECS)} %) are completely clean: canonical
`budget_id` and `client_code`, ISO dates, `EUR`, numeric amounts, no sentinels,
and `total_amount == hours_estimated × {RATE_EUR_PER_HOUR:.2f}` with `phases[]` summing to both.
The table below is exhaustive: every defect in the corpus is listed here, and
every row below is actually present on disk — the generator applies these rows
rather than describing them.

| File (under `data/corpus/`) | Field | Family | Planted value | Notes |
|---|---|---|---|---|
{chr(10).join(table_rows)}

### Count per family

| Family | Planted instances |
|---|---|
{chr(10).join(f"| {family} | {count} |" for family, count in sorted(family_counts.items()))}

### Reading the families

- **{F1}** — the same fact written several ways. Dates in three formats
  (`2024-03-15` in the clean majority, `15/03/2024`, `Mar 15 2024` — all three
  denote the same day). Currency as `EUR`, `eur`, `€`, `euros`. `ACME Corp.`
  (canonical, `budget_003.json`) also appearing as `Acme Corp`, `acme corp` and
  `ACME`. `total_amount` as a number in the clean files, as `"80000"` and as
  `"80.000,00"` in the heterogeneous ones.
- **{F2}** — two `budget_id` values appear in two files each with different
  `total_amount` *and* different `signed_at`, so a "keep the latest" rule has
  something to resolve: `BUDGET-2024-0312` (`budget_022.json` 48750.0 @ 2024-03-12
  vs `budget_023.json` 52500.0 @ 2024-05-02) and `BUDGET-2023-0188`
  (`budget_024.json` 31250.0 @ 2023-06-04 vs `budget_025.json` 29875.0 @ 2023-09-18).
  `budget_022.json` and `budget_024.json` carry no defect of their own and are
  listed here only as the counterparts of the collision — but because they
  participate in one, they are not counted in the clean majority above.
- **{F3}** — all seven sentinels appear: `"N/A"`, `"-"`, `"unknown"`, `"TBD"`,
  `"pendiente"`, `""` and `" "`, spread across `client_name`, `contact_email`
  and `account_manager`, plus a transcript whose attendee block is filled with
  the same placeholders.
- **{F4}** — a negative `total_amount`, an absurd one (99 000 000), a
  `hours_estimated` in the millions, a `signed_at` in the future, and one
  `budget_id` outside the canonical pattern.

## What is *not* in here

No parser, no FastAPI endpoint, no validation logic. This directory only
produces data; the code that consumes it lives elsewhere.
"""
    return _write(DATA_DIR / "README.md", content.encode("utf-8"))


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------


def main() -> None:
    """Rebuild the whole corpus from scratch and print a short manifest."""
    # Sanity-check the phase weights before writing anything: a template that
    # does not sum to 1.0 would silently produce budgets whose phases do not
    # add up to their hours, which looks exactly like a planted defect.
    for project_type, weights in PHASE_WEIGHTS.items():
        total_weight = round(sum(weight for _, weight in weights), 6)
        if total_weight != 1.0:
            raise ValueError(f"phase weights for {project_type} sum to {total_weight}, not 1.0")

    grid = [rate for _, *rates in RATE_CARD_ROWS for rate in rates]
    mean_rate = round(sum(grid) / len(grid), 2)
    if mean_rate != RATE_EUR_PER_HOUR:
        raise ValueError(f"rate card mean is {mean_rate}, not the house rate {RATE_EUR_PER_HOUR}")

    if CORPUS_DIR.exists():
        shutil.rmtree(CORPUS_DIR)

    groups = {
        "historical_budgets": build_budgets(),
        "meeting_transcripts": build_transcripts(),
        "proposal_templates": build_proposals(),
        "signed_contracts": build_contracts(),
        "rate_card": [build_rate_card()],
    }
    readme = write_readme()

    print(f"corpus -> {CORPUS_DIR}")
    grand_total = 0
    for name, paths in groups.items():
        size = sum(path.stat().st_size for path in paths)
        grand_total += size
        print(f"  {name:<20} {len(paths):>3} files  {size / 1024:8.1f} KB")
    print(f"  {'TOTAL':<20} {sum(len(p) for p in groups.values()):>3} files  {grand_total / 1024:8.1f} KB")
    print(f"readme -> {readme} ({len(PLANTED_DEFECTS)} planted defects documented)")


if __name__ == "__main__":
    main()
