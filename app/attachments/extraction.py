"""Attachment processing, chosen over uploading
files directly to a provider's Files API so the estimator stays independent
of whichever model the Router happens to be using for a given call (see
``llm_wrapper.py`` for why that fallback guarantee matters here).

Extraction is a self-validating pipeline, not a single best-effort call:

1. Try the cheapest local method for the file type (``pypdf`` for PDF,
   ``python-docx`` for Word).
2. Validate the result against hard rules, non-empty, a minimum amount of
   text per page, a minimum ratio of printable characters. These catch the
   two failure modes that actually matter: a scanned/image-only PDF (near-zero
   extracted text) and a garbled/mis-decoded extraction.
3. If validation fails, retry a PDF with a second local engine (``PyMuPDF``,
   which sometimes recovers text ``pypdf`` misses on unusual encodings).
4. If that still fails, report it as failed rather than silently returning
   garbage. The caller (``EstimationService``) then falls back to the LLM's
   own multimodal document support for that one file, the "no OCR available"
   case is exactly what that fallback exists for. This keeps the common case
   (born-digital PDFs, Word docs) fully local while still reaching an answer
   for the rare fully-scanned document, instead of failing the whole request.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

import structlog
from pypdf import PdfReader

log = structlog.get_logger()

MIN_CHARS_PER_PAGE = 30
MIN_PRINTABLE_RATIO = 0.85


@dataclass
class ExtractionResult:
    filename: str
    text: str
    method: str  # "pypdf" | "pymupdf" | "docx" | "failed"
    ok: bool
    note: str | None = None


def _printable_ratio(text: str) -> float:
    if not text:
        return 0.0
    printable = sum(1 for c in text if c.isprintable() or c in "\n\t")
    return printable / len(text)


def _validate(text: str, units: int) -> tuple[bool, str | None]:
    """Hard rules gate the result (fail → try the next method / give up).
    A soft rule is logged but never fails the extraction on its own."""
    stripped = text.strip()
    if not stripped:
        return False, "no text extracted"

    chars_per_unit = len(stripped) / max(units, 1)
    if chars_per_unit < MIN_CHARS_PER_PAGE:
        return False, f"only {chars_per_unit:.0f} chars/page, likely scanned or empty"

    ratio = _printable_ratio(stripped)
    if ratio < MIN_PRINTABLE_RATIO:
        return False, f"low printable-character ratio ({ratio:.2f}), likely garbled encoding"

    words = stripped.split()
    if words and (len(stripped) / len(words)) > 40:
        log.warning("attachment_extraction_soft_warning", reason="unusually long average word length")

    return True, None


def _extract_pdf_pypdf(raw: bytes) -> tuple[str, int]:
    reader = PdfReader(io.BytesIO(raw))
    parts = [f"--- Page {i} ---\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages, start=1)]
    return "\n\n".join(parts), len(reader.pages)


def _extract_pdf_pymupdf(raw: bytes) -> tuple[str, int]:
    import pymupdf  # imported lazily, only needed on the fallback path

    doc = pymupdf.open(stream=raw, filetype="pdf")
    try:
        parts = [f"--- Page {i} ---\n{page.get_text().strip()}" for i, page in enumerate(doc, start=1)]
        return "\n\n".join(parts), doc.page_count
    finally:
        doc.close()


def _extract_docx(raw: bytes) -> tuple[str, int]:
    from docx import Document  # python-docx, imported lazily, .docx only

    document = Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    # docx has no natural "page" concept, treat the whole document as one unit.
    return "\n".join(parts), 1


def _debug_attempt(filename: str, method: str, text: str, units: int, ok: bool, note: str | None) -> None:
    log.debug(
        "attachment_conversion_attempt",
        filename=filename,
        method=method,
        units=units,
        chars_extracted=len(text.strip()),
        valid=ok,
        note=note,
    )


def extract_attachment(filename: str, raw: bytes) -> ExtractionResult:
    name = filename.lower()
    log.debug("attachment_conversion_started", filename=filename, bytes=len(raw))
    try:
        if name.endswith(".pdf"):
            text, pages = _extract_pdf_pypdf(raw)
            ok, note = _validate(text, pages)
            _debug_attempt(filename, "pypdf", text, pages, ok, note)
            if ok:
                return ExtractionResult(filename, text, "pypdf", True)

            log.info("attachment_extraction_retry", filename=filename, from_method="pypdf", reason=note)
            text2, pages2 = _extract_pdf_pymupdf(raw)
            ok2, note2 = _validate(text2, pages2)
            _debug_attempt(filename, "pymupdf", text2, pages2, ok2, note2)
            if ok2:
                return ExtractionResult(filename, text2, "pymupdf", True)
            return ExtractionResult(filename, "", "failed", False, note=note2 or note)

        if name.endswith(".docx"):
            text, units = _extract_docx(raw)
            ok, note = _validate(text, units)
            _debug_attempt(filename, "docx", text, units, ok, note)
            if ok:
                return ExtractionResult(filename, text, "docx", True)
            return ExtractionResult(filename, "", "failed", False, note=note)

        log.debug("attachment_conversion_unsupported", filename=filename)
        return ExtractionResult(filename, "", "failed", False, note=f"unsupported file type: {filename}")
    except Exception as exc:  # noqa: BLE001, a corrupt upload must degrade, not crash the request
        log.warning(
            "attachment_extraction_error",
            filename=filename,
            error_type=type(exc).__name__,
            error=str(exc)[:200],
        )
        return ExtractionResult(filename, "", "failed", False, note=f"extraction error: {exc}")


def format_attachments_block(results: list[ExtractionResult]) -> str:
    """Render successfully-extracted attachments as the delimited blocks the
    prompt expects. Failed ones are omitted here, the caller decides whether
    to note the failure in text or retry via the LLM's multimodal fallback."""
    blocks = [
        f"<attachment filename='{result.filename}'>\n{result.text}\n</attachment>"
        for result in results
        if result.ok
    ]
    return "\n\n".join(blocks)
