"""Proposal-template DOCX parser.

DOCX is unusually cooperative: unlike PDF it carries explicit semantic
structure (paragraph styles, heading levels), so the document's own
hierarchy is readable rather than inferred from font sizes and positions.

The parser exploits that by emitting one unit **per section**, split on
headings, with the heading as ``section_title``. The alternative — one
unit per whole proposal — would mean a query about delivery timelines
retrieves an entire proposal and makes the model find the relevant part
inside it. Sections let retrieval return "the Cronograma section of
proposal X", which is both a better match and a citable one.
"""

from __future__ import annotations

import io

from app.ingest.models import ParsedUnit
from app.ingest.parsers.base import Parser


def _is_heading(paragraph) -> bool:
    style = (getattr(paragraph.style, "name", "") or "").lower()
    return style.startswith("heading") or style.startswith("título") or style.startswith("titulo")


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(row) + " |" for row in body]
    return "\n".join(lines)


class ProposalDocxParser:
    supported_formats = {"docx"}

    def parse(self, content: bytes, source_hint: str) -> list[ParsedUnit]:
        from docx import Document as DocxDocument  # lazy: heavy, docx-only path

        document = DocxDocument(io.BytesIO(content))
        doc_id = source_hint.rsplit("/", 1)[-1]
        title = document.core_properties.title or doc_id
        author = document.core_properties.author or None

        units: list[ParsedUnit] = []
        current_heading: str | None = None
        buffer: list[str] = []

        def flush() -> None:
            body = "\n".join(b for b in buffer if b.strip()).strip()
            if not body:
                return
            units.append(
                ParsedUnit(
                    content=(f"## {current_heading}\n\n{body}" if current_heading else body),
                    document_id=doc_id,
                    unit_key=f"section-{len(units) + 1:04d}",
                    document_title=title,
                    document_author=author,
                    section_title=current_heading,
                    extra={"section_index": len(units) + 1},
                )
            )

        for paragraph in document.paragraphs:
            if _is_heading(paragraph) and paragraph.text.strip():
                flush()
                buffer = []
                current_heading = paragraph.text.strip()
                continue
            buffer.append(paragraph.text)

        # Tables live outside the paragraph flow in python-docx's model, so
        # they are appended to the final section rather than placed
        # inline. Imprecise, and the honest trade: recovering true
        # document order needs the raw XML body, which is a large amount
        # of machinery for proposals whose tables sit at the end anyway.
        for table in document.tables:
            markdown = _table_to_markdown(table)
            if markdown:
                buffer.append("\n" + markdown)

        flush()
        return units


_: Parser = ProposalDocxParser()
